"""Generate reviewable Word from the current Markdown, without changing user originals."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

root = Path(__file__).resolve().parents[1]
doc = Document()
doc.styles['Normal'].font.name = 'Noto Sans CJK SC'
doc.styles['Normal'].font.size = Pt(11)
for name in ('TECHNICAL_SPEC.md', 'VALIDATION.md'):
    code = False
    for line in (root / name).read_text().splitlines():
        if line.startswith('```'):
            code = not code
            continue
        if not code and line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            doc.add_heading(line.lstrip('#').strip(), min(level, 3))
        else:
            p = doc.add_paragraph(line)
            if code:
                for run in p.runs: run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(9)
    doc.add_page_break()
path = root / 'OrangePi_Terminal_Technical_Spec.docx'
doc.save(path)
print(path)
